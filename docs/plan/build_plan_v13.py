# -*- coding: utf-8 -*-
"""生成《Demon 2.0 技术方案与实施计划书 v1.3》——反向集成：酒馆当 UI，我们当引擎。"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT_DIR = r"E:\dshwork\plan-doc"
OUT = os.path.join(OUT_DIR, "Demon-技术方案与实施计划书-v1.3.docx")
os.makedirs(OUT_DIR, exist_ok=True)

EA_FONT = "微软雅黑"
LATIN_FONT = "Segoe UI"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.4)
sec.right_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = LATIN_FONT
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(4)

for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    st = doc.styles[name]
    st.font.name = LATIN_FONT
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
    st.paragraph_format.space_after = Pt(6)


def heading(text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    set_run_font(run, size={0: 24, 1: 16, 2: 13, 3: 11.5}.get(level, 11), bold=True,
                 color=RGBColor(0x1F, 0x3A, 0x5F))
    return p


def para(text, size=10.5, bold=False, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    p.paragraph_format.space_after = Pt(2)


def numbered(text):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    p.paragraph_format.space_after = Pt(2)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(headers, rows, widths=None, font_size=8.5, header_fill="DCE6F1"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=font_size, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(hdr[i], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=font_size)
            p.paragraph_format.space_after = Pt(0)
    if widths:
        total = sum(widths) or 1
        scale = 16.0 / total
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w * scale)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def code_block(lines):
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), "Consolas")
        rfonts.set(qn("w:hAnsi"), "Consolas")
        rfonts.set(qn("w:eastAsia"), EA_FONT)
        run.font.size = Pt(8.5)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)


def footer_page_number():
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, size=9, color=RGBColor(0x66, 0x66, 0x66))


footer_page_number()

# ============================ 封面 ============================
for _ in range(5):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Demon")
set_run_font(run, size=34, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("技术方案与实施计划书 · v1.3")
set_run_font(run, size=18, color=RGBColor(0x33, 0x33, 0x33))

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("主方向：反向集成 —— 酒馆当界面，我们当引擎")
set_run_font(run, size=11, color=RGBColor(0x66, 0x66, 0x66))

for _ in range(7):
    doc.add_paragraph()

meta = [
    ("版本", "v1.3（主方向定稿，替代 v1.0~v1.2）"),
    ("日期", "2026-10-24"),
    ("团队", "2 人 · 合计约 80 小时/周"),
    ("1.0", "一周：酒馆连上我们的引擎，正常聊天 + 自动出声"),
    ("2.0", "两个月：Live2D 角色表现（Live2DViewerEX ExAPI）"),
    ("形态", "Qt C++ 原生后台程序（托盘/设置），不重写聊天界面"),
]
for k, v in meta:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{k}：{v}")
    set_run_font(run, size=10, color=RGBColor(0x44, 0x44, 0x44))
    p.paragraph_format.space_after = Pt(2)

doc.add_page_break()

# ============================ 1 方向 ============================
heading("1. 方向与核心结论", 1)
para("本版确定主方向为「反向集成」：不重写聊天界面，而是让 Demon 对外提供一个 OpenAI 兼容的 HTTP 端点，把酒馆（SillyTavern）当成前端来用。")
para("一句话分工：", bold=True)
bullet("酒馆负责：界面、角色卡、世界书、宏、prompt 拼装、历史管理——这些一律不重写。")
bullet("Demon 负责：模型路由（本地/云端）、语音合成与播放、Live2D 角色表现（2.0）、单片机和树莓派（3.0/4.0）。")

heading("1.1 为什么这样最划算", 2)
table(
    ["项目", "自研全套（v1.2 路线）", "反向集成（本版）"],
    [
        ["prompt / 世界书 / 宏 / 模板", "需复刻，18~23 人天", "0（酒馆拼好发给我们）"],
        ["聊天界面", "需自研 Qt 界面", "0（酒馆就是界面）"],
        ["酒馆插件", "可能需写", "0（走标准 OpenAI 兼容设置）"],
        ["1.0 工期", "约 80 小时", "**约 30 小时**"],
        ["升级风险", "需跟随酒馆数据格式", "只依赖一个稳定的设置项"],
        ["可控性", "完全自主", "受酒馆 UI 限制"],
    ],
    widths=[5.4, 6.4, 6.4],
)

para("代价也要说清楚：这个模式下我们不是“酒馆替身”，而是“酒馆的发动机”。哪天想脱离酒馆，再把 v1.2 里的数据层与推理管线补上、自研 UI——两条路不冲突，先走便宜的那条。")

# ============================ 2 架构 ============================
heading("2. 总体架构", 1)
code_block([
    "┌──────────────────────────── Windows PC ────────────────────────────┐",
    "│                                                                    │",
    "│  SillyTavern（浏览器 UI，用户照常用）                                │",
    "│      │ OpenAI 兼容请求 POST /v1/chat/completions (stream)           │",
    "│      ▼                                                             │",
    "│  Demon 引擎（Qt C++，常驻，127.0.0.1:8790）                          │",
    "│   ├─ HTTP/SSE 网关   /v1/models、/v1/chat/completions               │",
    "│   ├─ 上游路由        本地 llama.cpp / 云端 OpenAI 兼容 / 故障切换     │",
    "│   ├─ 后处理          去格式残留、情绪标签提取（2.0 用）                │",
    "│   ├─ 语音            调 api_v2(9880) 分句合成 → 播放                │",
    "│   └─ 事件总线        回复完成 → 通知 Live2D / 硬件                   │",
    "│                                                                    │",
    "│  （2.0）Live2DViewerEX  ◀──WebSocket ExAPI── 表情/动作/气泡/音频      │",
    "│  （3.0）STM32F103/G431 ◀──串口/BLE── 情绪/口型/按钮                  │",
    "└────────────────────────────────────────────────────────────────────┘",
])

heading("2.1 进程与启动", 2)
para("必须遵守的一条铁律：Demon 引擎要由用户或登录项独立启动，**不要从 Electron/浏览器里拉起**。上一阶段已经实测：在 Chromium 的进程树里 spawn 的 Python 会卡死在 650MB。")
table(
    ["方式", "可行性", "说明"],
    [
        ["登录项 / 快捷方式启动", "✅ 推荐", "Shell 血统，spawn 语音服务正常（实测 12~18 秒就绪）"],
        ["托盘手动启动", "✅ 推荐", "同上"],
        ["由酒馆/浏览器拉起", "❌ 禁止", "会继承 Chromium 环境，语音服务起不来"],
        ["由 Electron 外壳拉起", "❌ 禁止", "同上"],
    ],
    widths=[5.0, 2.6, 8.4],
)

# ============================ 3 对接契约 ============================
heading("3. 与酒馆的对接契约（1.0 的核心）", 1)
para("酒馆设置路径：Chat Completion → Chat Completion Source 选 Custom (OpenAI-compatible)，填端点与模型名。之后酒馆会把拼好的 prompt 发给我们，我们只需按 OpenAI 协议回话。")

heading("3.1 必须实现的端点", 2)
table(
    ["方法", "路径", "说明"],
    [
        ["GET", "/v1/models", "返回模型列表（至少一个，例如 demon），酒馆用它填充下拉框"],
        ["POST", "/v1/chat/completions", "核心：stream=true 时返回 SSE；false 时返回普通 JSON"],
        ["OPTIONS", "/v1/*", "CORS 预检，回 204 + 允许头"],
        ["GET", "/healthz", "运维用：引擎存活、上游连通性、语音服务状态"],
    ],
    widths=[1.8, 5.2, 9.0],
)

heading("3.2 请求/响应格式", 2)
code_block([
    "POST /v1/chat/completions",
    "{",
    '  "model": "demon",',
    '  "messages": [{"role":"system","content":"..."},{"role":"user","content":"..."}],',
    '  "stream": true,',
    '  "temperature": 1.0, "max_tokens": 400, "stop": ["..."],',
    '  ...其余字段一律忽略，不报错',
    "}",
    "",
    "流式响应（text/event-stream）：",
    'data: {"id":"x","object":"chat.completion.chunk","choices":[{"delta":{"content":"你"},"index":0}]}',
    'data: {"id":"x",...,"choices":[{"delta":{"content":"好"},"index":0}]}',
    "data: [DONE]",
])

heading("3.3 工程细节（不做就会踩坑）", 2)
table(
    ["项", "要求", "为什么"],
    [
        ["CORS", "Access-Control-Allow-Origin 回显请求 Origin；允许 Authorization/Content-Type", "酒馆页面是 http://127.0.0.1:8000，跨源请求"],
        ["预检 OPTIONS", "204 + Allow-Methods/Headers/Max-Age", "带 Authorization 的 POST 先发预检"],
        ["SSE 保活", "每 15 秒发一行 `:keep-alive`", "避免中间层/浏览器空闲超时断开"],
        ["结束标记", "必须发 `data: [DONE]`", "少了酒馆会一直转圈"],
        ["未知字段", "全部忽略，不返回 400", "不同酒馆版本会多带参数"],
        ["错误格式", "`{\"error\":{\"message\":\"...\"}}` + 4xx/5xx", "酒馆会把它显示成错误气泡"],
        ["取消", "客户端断开时中止上游请求", "否则白烧 token"],
        ["鉴权", "可选：校验 Custom API Key", "本机自用可关"],
    ],
    widths=[2.4, 8.0, 5.6],
)

heading("3.4 上游路由", 2)
bullet("本地：llama.cpp / LM Studio / vLLM 等 OpenAI 兼容服务（127.0.0.1:某端口）。")
bullet("云端：任意 OpenAI 兼容端点（DeepSeek、Moonshot、OpenAI 等），密钥存本地配置文件。")
bullet("路由策略：默认单上游；可选按模型名分流（如 demon-local / demon-cloud），失败自动回退。")
bullet("透传原则：请求体原样转发，只改 model 名；响应流原样回传，但顺带旁路解析出完整文本，供 TTS 与 Live2D 使用。")

# ============================ 4 1.0 ============================
heading("4. 1.0 规格（一周，约 30 小时）", 1)
heading("4.1 功能清单", 2)
table(
    ["模块", "内容", "估时"],
    [
        ["HTTP 网关", "/v1/models、/v1/chat/completions（SSE 透传）、OPTIONS、/healthz", "8h"],
        ["上游路由", "配置文件（本地/云端/密钥）、按模型名分流、超时与回退", "4h"],
        ["语音链路", "回复完成后分句 → api_v2 合成 → 播放队列；两条声线切换", "10h"],
        ["托盘与设置", "启停、开机自启开关、声线选择、日志窗口、打开配置目录", "4h"],
        ["健壮性", "日志落盘、异常不崩、上游断开重连、端口占用提示", "4h"],
    ],
    widths=[3.4, 10.0, 2.6],
)

heading("4.2 配置文件草案", 2)
code_block([
    "{",
    '  "listen": { "host": "127.0.0.1", "port": 8790 },',
    '  "upstreams": [',
    '    { "id": "local", "baseUrl": "http://127.0.0.1:8080/v1", "model": "qwen", "apiKey": "" },',
    '    { "id": "cloud", "baseUrl": "https://api.example.com/v1", "model": "chat", "apiKey": "sk-..." }',
    "  ],",
    '  "defaultUpstream": "local",',
    '  "voice": { "provider": "gptsovits", "endpoint": "http://127.0.0.1:9880",',
    '             "profile": "xiaod", "autoPlay": true, "speed": 1.0 },',
    '  "engine": { "keepAliveSeconds": 15, "requestTimeoutSeconds": 180 }',
    "}",
])

heading("4.3 验收清单", 2)
numbered("酒馆的下拉框里能看到 demon，并能选中。")
numbered("发一句话，酒馆里能流式出字（首字延迟受上游影响，但不出现整段等待）。")
numbered("回复结束后自动播报，音色是选定声线；关闭自动播报则静默。")
numbered("两条声线切换后，下一次播报音色随之改变。")
numbered("关掉 Demon 后酒馆给出明确报错（不是转圈）；重新启动引擎后无需重启酒馆即可恢复。")
numbered("引擎日志完整（请求、上游、TTS、错误），可直接排障。")
numbered("连续对话 20 轮无内存增长、无残留进程。")

# ============================ 5 2.0 ============================
heading("5. 2.0：Live2D 角色表现（ExAPI）", 1)
para("Live2DViewerEX 提供 ExAPI（WebSocket + JSON，格式 {msg, msgId, data}），官方文档：live2d.pavostudio.com/doc/zh-cn/exapi/。")
table(
    ["消息码", "能力", "本项目用法"],
    [
        ["10000 / 10002", "注册事件监听 / 接收点击", "点角色 → 触发对话或打开酒馆"],
        ["11000", "显示气泡文本（可带按钮并返回选项）", "短句气泡、选项式交互"],
        ["13000 / 13100", "设置 / 移除模型", "切换角色"],
        ["13200", "触发动作（group:motion）", "情绪 → 动作组"],
        ["13300 / 13302", "设置 / 清除表情", "情绪 → 表情"],
        ["13500 / 13501", "播放 / 停止声音（文件或 base64，3 频道）", "**TTS 音频推给它**，口型由它跟音频"],
        ["12010 / 14000", "背景 / 特效（雨、雪）", "场景氛围（可选）"],
        ["13400", "设置模型位置", "多显示器摆位"],
    ],
    widths=[3.0, 7.4, 6.6],
)
para("情绪来源：回复文本里带标签（如 [emotion:shy]）或关键词启发式；映射表放在配置文件里，可随时调。", size=9.5, color=RGBColor(0x55, 0x55, 0x55))

# ============================ 6 后续 ============================
heading("6. 3.0 / 4.0 / 5.0 概览", 1)
table(
    ["版本", "目标", "关键点"],
    [
        ["3.0", "PC ↔ 单片机", "串口/BLE 协议：情绪、口型包络、按钮；STM32F103/G431 固件（朋友主导）"],
        ["4.0", "树莓派 Demon", "派上只做显示与交互；引擎仍在 PC，派用浏览器指向 PC:8790"],
        ["5.0", "三端联动", "PC 跑模型 / 派接云端 / 单片机执行；同一套协议"],
        ["6.0", "做减法", "体积、内存、启动时间"],
    ],
    widths=[1.6, 4.4, 10.0],
)

# ============================ 7 排期 ============================
heading("7. 1.0 排期（7 天 / 2 人）", 1)
table(
    ["天", "A", "B", "交付"],
    [
        ["D1", "HTTP 骨架（/v1/models、/healthz、CORS）", "配置加载、日志框架", "curl 能通"],
        ["D2", "SSE 透传（含 [DONE]、保活、取消）", "上游 HTTP 客户端 + 错误映射", "酒馆能选中并出字"],
        ["D3", "TTS 客户端（分句、合成、缓存）", "播放队列与声线切换", "能出声"],
        ["D4", "回复完成事件 → 触发播报", "托盘菜单（启停/声线/日志）", "端到端可用"],
        ["D5", "异常路径（上游断、TTS 断、端口占用）", "设置持久化", "稳定性"],
        ["D6", "性能与内存回归（20 轮）", "打包（CMake + windeployqt）", "可交付版本"],
        ["D7", "文档、验收清单逐条过", "可选：Live2D 连接预研", "1.0 完成"],
    ],
    widths=[1.2, 5.6, 5.6, 3.6],
)

# ============================ 8 风险 ============================
heading("8. 风险与应对", 1)
table(
    ["风险", "影响", "概率", "应对"],
    [
        ["酒馆版本改动 Custom 端点行为", "连不上", "低", "锁酒馆版本；升级前先在测试目录验证"],
        ["SSE 格式细节不符", "一直转圈", "中", "按 3.2 严格实现；用 curl 与酒馆双验证"],
        ["TTS 首句延迟 2~7 秒", "体感卡", "高", "分句流水线 + 预取下一句 + 常用句缓存"],
        ["语音服务未启动", "无声", "中", "托盘显示状态；一键拉起（Shell 血统）"],
        ["Live2DViewerEX ExAPI 版本差异", "2.0 返工", "中", "先写探测脚本打印能力，再做适配层"],
        ["口型来源不确定（系统音频/麦克风）", "口型对不上", "中", "2.0 首日验证；不行就 13500 推音频让它内部驱动"],
        ["过度依赖酒馆", "无法独立", "中", "引擎与 UI 解耦；未来可接自研前端"],
    ],
    widths=[6.2, 2.8, 1.2, 6.4],
)

# ============================ 9 备选 ============================
heading("9. 备选路线（如果将来要脱离酒馆）", 1)
para("v1.2 里的两章仍然有效，作为“自研 UI”阶段的资产保留：")
bullet("酒馆兼容数据层（角色卡 PNG tEXt / jsonl / 世界书 / settings 子集），约 8~10 人天。")
bullet("推理管线（宏、世界书注入、历史预算、模板、流式、重生成），约 18~23 人天。")
bullet("两者都只在“要自己做界面”时才需要；当前路线不投入。")

heading("9.1 酒馆数据格式速查（备查）", 2)
table(
    ["文件", "关键字段"],
    [
        ["角色卡 JSON V2", "data{name,description,personality,scenario,first_mes,mes_example,creator_notes,system_prompt,post_history_instructions,alternate_greetings,character_book,tags,creator,character_version,extensions}"],
        ["PNG 卡", "tEXt 块 keyword=chara，值为 base64(UTF-8 JSON)"],
        ["聊天 jsonl", "首行含 user_name/character_name/create_date/chat_metadata；消息行含 name/is_user/is_system/send_date/mes/extra"],
        ["世界书", "entries{uid:{key,keysecondary,selective,selectiveLogic,order,position,depth,role,probability,disable,group,...}}"],
    ],
    widths=[3.6, 12.4],
)

# ============================ 10 下一步 ============================
heading("10. 下一步", 1)
numbered("确认 Qt 版本与编译器（建议 Qt 6.8+ 与 MSVC 2022；MinGW 也可但不推荐）。")
numbered("建仓库 Demon2/{engine,tray,live2d,mcu,docs}。")
numbered("先写 /healthz 与 /v1/models，用 curl 打通，再进 SSE。")
numbered("准备两个上游：本地 llama.cpp 与一个云端 OpenAI 兼容服务，做故障切换测试。")

# ============================ 附录 ============================
heading("附录：术语", 1)
table(
    ["术语", "含义"],
    [
        ["反向集成", "让既有 UI（酒馆）调用我们的服务，而不是我们重写 UI"],
        ["OpenAI 兼容端点", "按 OpenAI Chat Completions 协议提供 HTTP 服务"],
        ["SSE", "Server-Sent Events，服务端单向流式推送"],
        ["ExAPI", "Live2DViewerEX 提供给第三方程序的 WebSocket 控制接口"],
        ["上游", "真正做推理的服务（本地模型或云端 API）"],
    ],
    widths=[3.6, 12.4],
)

doc.save(OUT)
print("SAVED", OUT)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
