# -*- coding: utf-8 -*-
"""生成 Demon 技术方案与实施计划书 v1.4（聚焦 1.1 缺陷修复与 2.0 Live2D）。"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = r"E:\dshwork\plan-doc\Demon-技术方案与实施计划书-v1.4.docx"

doc = Document()

# 基础样式
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

for level in range(1, 4):
    h = doc.styles[f"Heading {level}"]
    h.font.name = "微软雅黑"
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
        shade(hdr[i], "E8ECF5")
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(text))
            run.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows:
                r.cells[i].width = Cm(w)
    return t


def bullet(text, level=0):
    p = doc.add_paragraph(text, style="List Bullet" if level == 0 else "List Bullet 2")
    return p


# ---------------- 封面 ----------------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Demon 技术方案与实施计划书")
run.font.size = Pt(26)
run.bold = True
run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x4A)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("v1.4 · 1.1 缺陷收尾 + 2.0 Live2D（VTube Studio 单后端）")
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x77)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("2026-10-24 · 内部方案 · 40h/周 + 朋友 40h/周").font.size = Pt(9.5)

doc.add_paragraph()
table(
    ["项目", "内容"],
    [
        ["产品名", "Demon（不改名）"],
        ["当前版本", "1.1.1（千问语音 + Live2D 起始）"],
        ["本版目标", "修完 1.1 暴露的问题，2.0 完成 VTube Studio 单后端交互"],
        ["明确不做", "单片机 / 树莓派（4.0）、MCU（3.0）、多引擎（5.0）"],
        ["运行环境", "Windows 10/11 x64，Qt 6.8.3 + MSVC 2022"],
    ],
    widths=[3.2, 12.5],
)

doc.add_page_break()

# ---------------- 1 版本定位 ----------------
doc.add_heading("1. 版本定位与节奏", level=1)
doc.add_paragraph(
    "1.1 的目标是“能稳定用起来”：语音换成千问云端、删掉未达预期的自动播报、"
    "补上 Live2D 的基础接线。实际使用后暴露了四类问题：音频输出不受控、音色不是参考音频、"
    "启动页白底、酒馆保存报错。这些问题必须在 2.0 之前收干净，否则 2.0 的联调会被杂音干扰。"
)
table(
    ["版本", "定位", "出口标准"],
    [
        ["1.1.x", "缺陷修复", "四类问题全部关闭，回归清单通过"],
        ["2.0", "VTube Studio 单后端", "表情 / 口型 / 动作 / 点击 四件事可用"],
        ["3.0", "PC ↔ 单片机", "本版不展开"],
    ],
    widths=[2.2, 6.5, 7.0],
)
doc.add_paragraph(
    "排期按每周 40h（本人）+ 40h（朋友）估算，2.0 的 MVP 目标 4 周内可验收。"
)

# ---------------- 2 1.1 遗留问题 ----------------
doc.add_heading("2. 1.1 遗留问题与修复方案", level=1)

doc.add_heading("2.1 音频输出绕过耳机与系统音量", level=2)
doc.add_paragraph(
    "现象：点“试听千问”时声音从笔记本扬声器外放，耳机没有声音；系统音量滑杆对其无效。"
)
doc.add_paragraph("原因分析：")
bullet("QAudioOutput 在创建时绑定当时的默认设备，之后耳机插拔/默认设备变化不会跟随。")
bullet("播放器音量固定为 1.0（代码里写死），没有暴露给用户，也没有和 config 的 voice.volume 打通。")
doc.add_paragraph("修复：")
bullet("每次播放前对比 QMediaDevices::defaultAudioOutput()，设备变了就重设，保证“跟随系统默认”。")
bullet("语音页新增“输出”分组：设备下拉（系统默认 + 枚举所有输出设备）+ 音量滑杆，写入 config。")
bullet("配置项 voice.outputDevice（空 = 系统默认）、voice.volume（0~1）。")
table(
    ["项", "状态", "验证方式"],
    [
        ["跟随系统默认设备", "已实现", "插拔耳机后试听，声音走耳机"],
        ["手动选设备", "已实现", "下拉选耳机，试听走耳机"],
        ["音量可调", "已实现", "滑杆 0/50/100 三档试听"],
    ],
    widths=[4.5, 2.5, 8.5],
)

doc.add_heading("2.2 千问音色不是参考音频", level=2)
doc.add_paragraph(
    "现象：千问合成用的是内置音色 longanhuan_v3.1，和 voices/ 里的参考音频（小d、demon）不是同一个声音。"
)
doc.add_paragraph("结论：这不是 bug，是缺功能。Qwen-Audio-TTS 的内置音色无法“听起来像参考音频”，"
                  "必须走声音复刻（Voice Cloning）：用参考音频创建一个音色 ID，再用该 ID 合成。")
doc.add_paragraph("方案（2.0 期间并行做）：")
table(
    ["步骤", "做法", "产出"],
    [
        ["1. 准备样本", "从 voices/*.wav 选 10~60 秒干净人声，单声道 16k/24k", "ref_clean.wav"],
        ["2. 创建音色", "调 Qwen 声音复刻接口上传样本，得到 voice_id", "voice_id 写入 config"],
        ["3. 合成", "voice.qwen.voice = <voice_id>，其余不变", "音色接近参考音频"],
        ["4. 校验", "试听 + 与参考音频做主观对比", "验收记录"],
    ],
    widths=[3.0, 9.5, 3.5],
)
doc.add_paragraph(
    "注意：声音复刻涉及他人声音授权。仓库里的参考音频需确认来源与授权，"
    "发布包不内置任何音色样本。"
)

doc.add_heading("2.3 启动页变白", level=2)
doc.add_paragraph(
    "现象：启动页内容区整块白色，和深色主题不一致。"
)
doc.add_paragraph(
    "原因：启动页内容变长后套了 QScrollArea，滚动区默认使用系统浅色背景，且 viewport 自动填充。"
)
doc.add_paragraph(
    "修复：QScrollArea 及内层 widget 背景设为透明，viewport 关闭 autoFillBackground，"
    "并统一滚动条配色（#1b1b2b 底、#3a3a55 滑块）。"
)

doc.add_heading("2.4 酒馆报“聊天无法保存”", level=2)
doc.add_paragraph(
    "现象：在酒馆发消息后弹出“聊天无法保存，检查服务器连接并刷新页面以防止数据丢失”。"
)
doc.add_paragraph("定位过程与结论：")
bullet("磁盘充足（C: 剩 22.8 GB），chats 目录可写，服务在 8000 正常监听。")
bullet("发现 chats 根目录出现一个 0 字节的同名聊天文件（真正的聊天在 小鲸鱼娘\\ 子目录，571 KB）。"
       "这是页面与服务端会话不一致时，保存走了错误路径并写失败留下的残骸。")
bullet("酒馆服务在 23:35 重启过一次，浏览器页面仍是重启前的会话；CSRF/会话失效后保存请求被拒绝。")
doc.add_paragraph("处理：")
bullet("已删除 0 字节残留文件，避免酒馆把它当成一个空聊天。")
bullet("使用侧：刷新酒馆页面（F5）后重新进入角色聊天；如果仍报错，关掉酒馆与页面，用 Demon 启动器重新拉起。")
bullet("代码侧（2.0 一并做）：Demon 启动器检测到 8000 已监听时只“打开页面”，不重复启动，减少服务被重启的次数。")
table(
    ["检查项", "结果"],
    [
        ["C: 剩余空间", "22.8 GB，正常"],
        ["chats 目录写入", "可写"],
        ["0 字节残留", "已清理"],
        ["服务端口", "8000 正常"],
        ["页面会话", "需刷新后重新建立"],
    ],
    widths=[5.5, 10.0],
)

doc.add_heading("2.5 其他遗留项", level=2)
table(
    ["编号", "问题", "处理"],
    [
        ["B-01", "Demon 以管理员运行时，普通权限无法结束进程", "用户自行选择一次性管理员；文档标注"],
        ["B-02", "托盘图标在 Win11 上 geometry 为 0", "已有诊断日志，控制面板为主入口"],
        ["B-03", "日志无限增长", "已做 8 MB 滚动，保留 1 份备份"],
        ["B-04", "双实例", "QSharedMemory 单实例锁，已生效"],
    ],
    widths=[2.0, 8.5, 5.5],
)

# ---------------- 3 2.0 目标 ----------------
doc.add_heading("3. 2.0 目标：VTube Studio 单后端", level=1)
doc.add_paragraph(
    "2.0 只做一件事：让 Demon 稳定地驱动 VTube Studio 里的 Live2D 模型。"
    "不引入第二套渲染器，不在 Demon 内嵌 Live2D 渲染，所有画面交给 VTS。"
)
doc.add_heading("3.1 范围", level=2)
table(
    ["能力", "说明", "优先级"],
    [
        ["连接与授权", "WebSocket 8001，token 持久化，断线重连", "P0"],
        ["表情", "按情绪激活 exp3，支持多表情叠加", "P0"],
        ["口型", "TTS 音频包络 → MouthOpen（30fps）", "P0"],
        ["水印常驻", "大肥鱼模型自动激活“13 水印消除”", "P0"],
        ["动作", "TriggerHotkey 触发热键动作", "P1"],
        ["点击交互", "ModelClickedEvent → 气泡/回应", "P1"],
        ["情绪识别", "文本关键词 → 情绪 → 表情", "P1"],
        ["表情复位", "N 秒无操作回到中性", "P2"],
    ],
    widths=[3.0, 9.5, 3.0],
)
doc.add_heading("3.2 非目标", level=2)
bullet("不做 Live2D 自渲染（Cubism SDK / pixi）。")
bullet("不做面捕、不做模型编辑。")
bullet("不做第二后端（Live2DViewerEX 留作备选，不排期）。")

# ---------------- 4 架构 ----------------
doc.add_heading("4. 架构与接口", level=1)
doc.add_paragraph(
    "数据流：酒馆 → Demon(/v1/chat/completions) → 上游模型；"
    "插件按句请求 /voice/speak → 千问合成 → 返回 WAV 并驱动口型；"
    "情绪由文本或标签触发 → VTS 表情。"
)
doc.add_heading("4.1 组件", level=2)
table(
    ["组件", "职责"],
    [
        ["AppController", "HTTP 路由、语音编排、Live2D 编排、启动器"],
        ["TtsClient", "gptsovits / qwen 两种合成后端"],
        ["VtsClient", "VTS WebSocket：认证、参数注入、表情、热键"],
        ["LipSync", "WAV → 30fps 包络"],
        ["MainWindow", "状态 / 接口 / 启动 / 语音 / 角色 / Live2D / 日志"],
    ],
    widths=[4.0, 11.5],
)
doc.add_heading("4.2 HTTP 接口", level=2)
table(
    ["方法", "路径", "说明"],
    [
        ["GET", "/healthz", "版本、上游、语音概要"],
        ["POST", "/v1/chat/completions", "SSE 透传"],
        ["POST", "/voice/speak", "合成 WAV，可带 emotion"],
        ["GET", "/voice/status", "声线、后端、设备"],
        ["GET", "/live2d/status", "VTS 连接状态"],
        ["POST", "/live2d/emotion", "切换情绪表情"],
        ["POST", "/live2d/test", "注入 MouthOpen"],
        ["GET", "/live2d/expressions", "当前模型表情表"],
        ["POST", "/live2d/raw", "调试：任意 VTS 请求"],
    ],
    widths=[2.0, 6.0, 7.5],
)
doc.add_heading("4.3 配置结构（节选）", level=2)
code = doc.add_paragraph()
code_run = code.add_run(
    '{\n'
    '  "voice": {\n'
    '    "provider": "qwen",\n'
    '    "outputDevice": "",\n'
    '    "volume": 1.0,\n'
    '    "qwen": { "model": "qwen-audio-3.1-tts-flash", "voice": "longanhuan_v3.1" }\n'
    '  },\n'
    '  "live2d": {\n'
    '    "host": "127.0.0.1", "port": 8001,\n'
    '    "vtsExe": "E:/.../VTube Studio.exe",\n'
    '    "modelsDir": "E:/.../Live2DModels"\n'
    '  }\n'
    '}'
)
code_run.font.name = "Consolas"
code_run.font.size = Pt(9)

# ---------------- 5 关键实现 ----------------
doc.add_heading("5. 关键实现与已知坑", level=1)
doc.add_heading("5.1 认证", level=2)
doc.add_paragraph(
    "连接后先发 AuthenticationRequest；没有 token 时发 AuthenticationTokenRequest，"
    "VTS 弹窗由用户点“允许”，token 存 data/vts-token.json。"
)
doc.add_heading("5.2 表情协议（实测）", level=2)
table(
    ["项", "结论"],
    [
        ["消息类型", "ExpressionActivationRequest"],
        ["字段", "expressionFile（文件名）+ active: true"],
        ["是否叠加", "可以，多个表情同时生效"],
        ["参数注入", "只能注入 tracking 参数，Paramcol_3 不可注入"],
    ],
    widths=[4.0, 11.5],
)
doc.add_heading("5.3 水印常驻", level=2)
doc.add_paragraph(
    "大肥鱼的 Paramcol_3 不能直接注入，改为激活 EXP/13 水印消除.exp3.json。"
    "切到该模型时激活一次，之后每次切情绪前补一次，保证叠加后仍在。"
)
doc.add_heading("5.4 口型", level=2)
doc.add_paragraph(
    "LipSync 解析 WAV（PCM16），按 30fps 窗口算 RMS，归一化到 0~1；"
    "播放时用 QTimer 逐帧 InjectParameterDataRequest 注入 MouthOpen，结束后归零。"
)
doc.add_heading("5.5 VTS 兼容性坑", level=2)
table(
    ["坑", "表现", "对策"],
    [
        ["ExpressionStateResponse 非法 UTF-8", "c_0120 上收到 1007 断线", "本地读 model3.json 表情表"],
        ["事件订阅 ModelLoadedEvent", "连接被关闭", "改为 15 秒轮询 CurrentModelRequest"],
        ["模型切换期间请求", "VTS 关连接", "模型响应后延迟 400ms 再取表情"],
        ["重连风暴", "VTS 拒绝新连接", "重连上限 12 次，之后转人工"],
    ],
    widths=[5.0, 6.0, 5.0],
)
doc.add_heading("5.6 情绪映射", level=2)
doc.add_paragraph("两套模型各一张表，键为情绪 ID（happy/shy/angry/...），值为表情名关键字。")
table(
    ["情绪", "大肥鱼", "c_0120"],
    [
        ["happy", "02 龇牙", "开心兴奋"],
        ["shy / love", "04 脸红", "脸红 / 爱心眼"],
        ["sad / cry / tired", "08 担心嘴", "悲伤 / 哭 / 流汗"],
        ["angry / dark", "03 脸黑", "生气 / 阴暗"],
        ["surprised", "05 O O", "感叹号"],
        ["confused / dizzy", "06 X X", "问号 / 晕晕"],
        ["think", "08 栗子嘴", "呆呆眼"],
        ["tease", "05 》《", "调皮"],
    ],
    widths=[4.0, 6.0, 6.0],
)

# ---------------- 6 里程碑 ----------------
doc.add_heading("6. 里程碑与排期", level=1)
table(
    ["阶段", "内容", "工期", "出口"],
    [
        ["M1（本周）", "1.1 四个问题收尾 + 回归", "2 天", "清单全绿"],
        ["M2（第 2 周）", "表情/口型/水印稳定 + 两模型映射", "1 周", "连续 30 分钟不假死"],
        ["M3（第 3 周）", "动作热键 + 点击交互", "1 周", "点模型有反应"],
        ["M4（第 4 周）", "情绪复位 + 打磨 + 1.1.2 打包", "3 天", "验收清单通过"],
    ],
    widths=[2.5, 8.5, 2.5, 4.0],
)

# ---------------- 7 验收 ----------------
doc.add_heading("7. 验收清单", level=1)
checks = [
    "插拔耳机后试听，声音走当前默认设备；音量滑杆生效",
    "千问试听音色与参考音频接近（复刻完成后）",
    "启动页深色，无白底",
    "酒馆发消息不再弹“聊天无法保存”",
    "VTS 启动/关闭可从 Demon 启动器完成",
    "大肥鱼：切换情绪后水印仍在",
    "c_0120：44 个表情可加载，情绪切换正确",
    "说话时口型跟随，结束后归零",
    "点击模型有反馈（气泡或动作）",
    "连续运行 30 分钟，VTS 连接不断",
]
for c in checks:
    doc.add_paragraph(c, style="List Number")

# ---------------- 8 风险 ----------------
doc.add_heading("8. 风险与对策", level=1)
table(
    ["风险", "影响", "对策"],
    [
        ["VTS 版本升级改协议", "连接/表情失效", "所有调用集中 VtsClient，出问题只改一处"],
        ["声音复刻授权", "不能公开分发音色", "样本不打包，用户自备"],
        ["模型许可", "不能商用/再分发", "模型目录附 README 与许可说明"],
        ["管理员权限差异", "联调环境不一致", "开发用管理员，验收用普通权限各跑一遍"],
    ],
    widths=[4.5, 5.5, 6.0],
)

doc.add_paragraph()
tail = doc.add_paragraph("—— 本方案随实现更新，变更记录见 docs/DEVLOG.md ——")
tail.alignment = WD_ALIGN_PARAGRAPH.CENTER
tail.runs[0].font.size = Pt(9)

doc.save(OUT)
print("saved:", OUT)
