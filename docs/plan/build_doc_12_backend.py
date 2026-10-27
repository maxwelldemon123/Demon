# -*- coding: utf-8 -*-
"""生成 Demon 1.2 前后端数据结构与后端实现说明（Word）。"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = r"E:\dshwork\plan-doc\Demon-1.2-前后端数据结构与实现说明.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
for level in range(1, 4):
    h = doc.styles[f"Heading {level}"]
    h.font.name = "微软雅黑"
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(text))
            run.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows:
                r.cells[i].width = Cm(w)
    return t


def code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    return p


title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Demon 1.2 前后端数据结构与后端实现说明")
run.font.size = Pt(22)
run.bold = True
run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x4A)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("2026-10-26 · 供优化方向评审用").font.size = Pt(9.5)

doc.add_heading("1. 总体结构", level=1)
doc.add_paragraph(
    "Demon 是单进程 Qt Widgets 应用：主线程跑 UI 与 HTTP 服务，网络请求走 QNetworkAccessManager 的异步回调，"
    "无独立工作线程。各模块通过 AppController 汇总，UI 只读写 AppController。"
)
table(
    ["层", "组件", "职责"],
    [
        ["UI", "MainWindow", "侧栏导航 + 7 个页面；只调用 AppController 的公开方法与信号"],
        ["控制", "AppController", "路由分发、语音编排、Live2D 编排、启动器、配置读写"],
        ["HTTP 服务", "HttpServer", "手写 QTcpServer，HTTP/1.1 解析、SSE 流式响应"],
        ["语音", "TtsClient / AudioPlayer / LipSync", "本地 GPT-SoVITS、千问/CosyVoice 合成、播放、口型包络"],
        ["Live2D", "VtsClient", "VTube Studio WebSocket 协议：认证、参数注入、表情、热键"],
        ["配置", "Config", "JSON 读写，全部字段默认值"],
    ],
    widths=[2.4, 5.2, 8.0],
)

doc.add_heading("2. 配置数据结构（config.json）", level=1)
doc.add_paragraph("Config 是唯一持久化状态；UI 修改 m_config 后调用 saveConfig()，需要热重载时调用 reloadConfig()。")
doc.add_heading("2.1 顶层", level=2)
code('''{
  "listen":   { "host": "127.0.0.1", "port": 8790 },
  "modelId":  "demon",
  "defaultUpstream": "local",
  "upstreams": [ UpstreamConfig ],
  "voice":     VoiceConfig,
  "voiceProfiles": [ VoiceProfileConfig ],
  "characterVoices": { "角色名": "声线档案 id" },
  "charactersDir": "",
  "launcher":  LauncherConfig,
  "live2d":    Live2dConfig,
  "engine":    { "keepAliveSeconds": 15 }
}''')
doc.add_heading("2.2 子结构", level=2)
table(
    ["结构", "字段", "说明"],
    [
        ["UpstreamConfig", "id / baseUrl / model / apiKey", "OpenAI 兼容上游；model 会改写请求体"],
        ["VoiceConfig", "provider / endpoint / profile / speed / volume / outputDevice / translateToJa / qwen", "provider: gptsovits 或 qwen；输出设备空=系统默认"],
        ["QwenTtsConfig", "baseUrl / apiKey / model / voice / format / sampleRate", "model 可为 qwen-audio-3.1-tts-flash 或 cosyvoice-v3.5-flash"],
        ["VoiceProfileConfig", "id / name / gptWeights / sovitsWeights / refAudio / promptText / promptLang / qwenVoice", "qwenVoice = 复刻音色 ID"],
        ["Live2dConfig", "host / port / vtsExe / modelsDir", "VTS WebSocket 与模型目录"],
        ["LauncherConfig", "stPath / stPort / gsvPath", "酒馆与 GPT-SoVITS 路径"],
    ],
    widths=[3.0, 7.0, 5.5],
)

doc.add_heading("3. HTTP 接口契约", level=1)
doc.add_paragraph("统一 JSON；错误为 {\"error\":{\"message\":\"...\"}}。SSE 接口直接透传上游字节流。")
table(
    ["方法", "路径", "请求", "响应"],
    [
        ["GET", "/healthz", "—", "{status, version, model, upstreams, voice}"],
        ["GET", "/v1/models", "—", "OpenAI model list"],
        ["POST", "/v1/chat/completions", "OpenAI chat 请求", "stream=true 时 SSE；否则 501"],
        ["POST", "/voice/speak", "{text, profile, lang, speed?, emotion?}", "audio/wav 二进制"],
        ["GET", "/voice/status", "—", "{voice:{provider,model,profile,profiles[]}}"],
        ["GET", "/live2d/status", "—", "{live2d:{state,host,port}}"],
        ["POST", "/live2d/emotion", "{emotion}", "{emotion,model,applied}"],
        ["POST", "/live2d/test", "{param,value}", "{sent,param,value,state}"],
        ["GET", "/live2d/expressions", "—", "{model,expressions:[{name,file}]}"],
        ["POST", "/live2d/raw", "{type,data}", "{sent:true}（调试）"],
        ["GET", "/launcher/open", "—", "{opened:url}"],
    ],
    widths=[1.5, 5.0, 6.0, 5.0],
)

doc.add_heading("4. 语音链路", level=1)
doc.add_heading("4.1 合成", level=2)
doc.add_paragraph(
    "TtsClient 根据 provider 分流：gptsovits 走 api_v2 /tts；qwen 走 DashScope SpeechSynthesizer。"
    "qwen 分支返回的是 24 小时 URL，Demon 再下载成 WAV 回给酒馆。"
)
code('''POST /voice/speak
  → profile = voiceProfiles 中 id 匹配项
  → provider=qwen:
      body = { model: voice.qwen.model,
               input: { text, voice: profile.qwenVoice ?: voice.qwen.voice,
                        format:"wav", sample_rate:24000, rate:speed } }
      → 响应 output.audio.url → GET → wav
  → 返回 wav，同时 LipSync 计算 30fps 包络推 MouthOpen''')
doc.add_heading("4.2 中文→日语", level=2)
doc.add_paragraph(
    "由 smartdemon 插件在前端完成：输出语言非 auto 且与文本语言不同 → 调用酒馆 generateQuietPrompt 翻译 → 再请求 /voice/speak。"
    "Demon 侧另有 translateToJa 兜底（走默认上游），默认关闭。"
)
doc.add_heading("4.3 播放", level=2)
doc.add_paragraph(
    "酒馆插件拿到 WAV 后用 <audio> 播放；Demon 自身 AudioPlayer 只用于「试听」，"
    "QAudioOutput 绑定 QMediaDevices 默认设备（可配置）。"
)

doc.add_heading("5. Live2D 链路", level=1)
table(
    ["阶段", "实现"],
    [
        ["连接", "ws://127.0.0.1:8001，AuthenticationRequest；无 token 时 AuthenticationTokenRequest，token 存 data/vts-token.json"],
        ["模型", "CurrentModelResponse 提供 modelName / modelID；AppController 用 modelID 匹配 .vtube.json 定位模型目录"],
        ["表情", "本地读 model3.json 的 Expressions；没有则递归扫 *.exp3.json；激活用 expressionFile + active"],
        ["叠加与清除", "表情可叠加；切情绪先 active:false 旧表情，8 秒无新情绪自动清空；水印表情常驻"],
        ["口型", "LipSync 解析 PCM16 WAV，30fps RMS 包络；QTimer 逐帧 InjectParameterData(MouthOpen)"],
        ["断线处理", "重连上限 12 次；ExpressionStateRequest 在部分模型会触发 1007，已改为本地读表情表"],
    ],
    widths=[3.0, 12.5],
)

doc.add_heading("6. UI 与后端的交互", level=1)
doc.add_paragraph(
    "MainWindow 构造时创建各页，把控件事件直接连到 AppController 方法；AppController 提供 uiDirty 信号，"
    "MainWindow 每 2 秒刷新状态栏/表格。"
)
table(
    ["UI 动作", "后端调用", "写配置"],
    [
        ["启动/停止酒馆", "startSillyTavern / stopSillyTavern", "否"],
        ["启动/停止语音服务", "startVoiceService / stopVoiceService", "否"],
        ["启动/关闭 VTS", "startVTubeStudio / stopVTubeStudio", "vtsExe 可写"],
        ["切换语音后端", "mutableConfig().voice + saveConfig + reloadConfig", "是"],
        ["选择输出设备/音量", "setOutputDevice / setVoiceVolume", "是"],
        ["声线表编辑", "saveVoiceTable（MainWindow）→ saveConfig + reloadConfig", "是"],
        ["角色绑定", "setCharacterVoice", "是"],
        ["情绪按钮", "testEmotion → applyEmotion", "否"],
        ["重载配置", "reloadConfig", "否"],
    ],
    widths=[4.5, 7.5, 3.5],
)

doc.add_heading("7. 已知问题与优化方向", level=1)
table(
    ["问题", "影响", "建议"],
    [
        ["AppController 单体过大（约 1500 行）", "UI 扩展后更难维护", "拆 VoiceService / LauncherService / Live2dService"],
        ["UI 直接改 m_config 再 saveConfig", "多处重复、易漏字段", "引入统一的 SettingsModel + 变更通知"],
        ["语音链路串行", "CosyVoice 每句 3~5 秒，播放有停顿", "预合成下一句 / 流式 TTS"],
        ["翻译在前端插件", "Demon 自己无法直接产出日语", "把翻译下沉到 Demon（可配置模型）"],
        ["表情表每次模型切换重扫目录", "大目录时首帧变慢", "缓存 modelID → 表情表"],
        ["日志只写文件", "UI 删除日志页后定位问题不便", "保留 /logs 接口或日志窗口"],
    ],
    widths=[5.0, 6.0, 5.0],
)

doc.add_paragraph()
tail = doc.add_paragraph("—— 供 1.2 优化方向讨论，数据结构以代码为准 ——")
tail.alignment = WD_ALIGN_PARAGRAPH.CENTER
tail.runs[0].font.size = Pt(9)

doc.save(OUT)
print("saved:", OUT)
